"""Build a blank life insurance tracking workbook.

Sheets:
  How to Use     - legend and instructions
  Policies       - one row per policy (inputs + calculated columns)
  Beneficiaries  - one row per beneficiary per policy
  Payments       - premium payment log
  Summary        - totals and counts driven by formulas
  Lists          - dropdown values (edit to customize)

Usage:
  python3 scripts/build_life_insurance_tracker.py [--output life_insurance_tracker.xlsx]
"""

import argparse
import datetime as dt

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = int(__import__("os").environ.get("TRACKER_ROWS", 300))  # pre-formatted rows per data sheet
FONT = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
EXAMPLE_FONT = Font(name=FONT, color="0000FF", italic=True)
INPUT_FONT = Font(name=FONT, color="0000FF")
CALC_FONT = Font(name=FONT, color="000000")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURRENCY = '$#,##0.00;($#,##0.00);"-"'
CURRENCY0 = '$#,##0;($#,##0);"-"'
DATE = "mm/dd/yyyy"
PCT = '0%;-0%;"-"'

LISTS = {
    "Policy Type": [
        "Term", "Whole Life", "Universal Life", "Indexed UL", "Variable UL",
        "Final Expense", "Group / Employer", "Other",
    ],
    "Status": [
        "Quoted", "Application Submitted", "Underwriting", "Approved",
        "In Force", "Lapsed", "Cancelled", "Declined", "Paid Out",
    ],
    "Frequency": ["Monthly", "Quarterly", "Semi-Annual", "Annual", "Single Pay"],
    "Payments / Year": [12, 4, 2, 1, 0],
    "Payment Method": ["Bank Draft", "Credit Card", "Check", "Payroll Deduction", "Other"],
    "Beneficiary Type": ["Primary", "Contingent"],
    "Relationship": [
        "Spouse", "Child", "Parent", "Sibling", "Other Relative",
        "Trust", "Estate", "Business", "Other",
    ],
}

# (header, width, kind, number_format)  kind: input | calc
POLICY_COLS = [
    ("Policy Number", 16, "input", None),
    ("Insured", 20, "input", None),
    ("Owner", 20, "input", None),
    ("Carrier", 22, "input", None),
    ("Policy Type", 16, "input", None),
    ("Status", 20, "input", None),
    ("Face Amount", 14, "input", CURRENCY0),
    ("Premium", 12, "input", CURRENCY),
    ("Frequency", 12, "input", None),
    ("Annual Premium", 14, "calc", CURRENCY),
    ("Application Date", 13, "input", DATE),
    ("Issue Date", 13, "input", DATE),
    ("Term Length (Yrs)", 10, "input", "0"),
    ("Term Ends", 13, "calc", DATE),
    ("Payment Method", 16, "input", None),
    ("Draft Day of Month", 10, "input", "0"),
    ("Next Premium Due", 13, "input", DATE),
    ("Days Until Due", 10, "calc", '0;"Overdue "0;"Today"'),
    ("Total Paid", 13, "calc", CURRENCY),
    ("Primary Benef. %", 10, "calc", PCT),
    ("Agent / Contact", 20, "input", None),
    ("Notes", 36, "input", None),
]

BENEF_COLS = [
    ("Policy Number", 16, "input", None),
    ("Beneficiary Name", 24, "input", None),
    ("Relationship", 16, "input", None),
    ("Type", 12, "input", None),
    ("Share %", 10, "input", PCT),
    ("Insured (lookup)", 20, "calc", None),
    ("Notes", 36, "input", None),
]

PAYMENT_COLS = [
    ("Date Paid", 13, "input", DATE),
    ("Policy Number", 16, "input", None),
    ("Amount", 12, "input", CURRENCY),
    ("Method", 16, "input", None),
    ("Insured (lookup)", 20, "calc", None),
    ("Notes", 36, "input", None),
]


def col(cols, header):
    return get_column_letter([c[0] for c in cols].index(header) + 1)


def write_headers(ws, cols):
    for i, (name, width, kind, _) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
        if kind == "calc":
            cell.comment = Comment("Calculated automatically - do not type here.", "Tracker")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2" if ws.title == "Policies" else "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ROWS + 1}"


def style_body(ws, cols):
    for r in range(2, ROWS + 2):
        for i, (_, _, kind, fmt) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            cell.border = BORDER
            cell.fill = CALC_FILL if kind == "calc" else INPUT_FILL
            cell.font = CALC_FONT if kind == "calc" else INPUT_FONT
            if fmt:
                cell.number_format = fmt


def list_ref(name):
    idx = list(LISTS).index(name) + 1
    letter = get_column_letter(idx)
    return f"Lists!${letter}$2:${letter}${len(LISTS[name]) + 1}"


def add_dropdown(ws, cols, header, list_name):
    c = col(cols, header)
    dv = DataValidation(type="list", formula1=f"={list_ref(list_name)}", allow_blank=True)
    dv.error = "Pick a value from the list (edit the Lists sheet to add options)."
    dv.errorTitle = "Invalid entry"
    ws.add_data_validation(dv)
    dv.add(f"{c}2:{c}{ROWS + 1}")


def build_lists(wb):
    ws = wb.create_sheet("Lists")
    for i, (name, values) in enumerate(LISTS.items(), start=1):
        h = ws.cell(row=1, column=i, value=name)
        h.font = Font(name=FONT, bold=True, color="FFFFFF")
        h.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = 22
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=i, value=v).font = INPUT_FONT
    ws.cell(row=len(LISTS["Status"]) + 4, column=1,
            value="Payments / Year lines up with Frequency and drives Annual Premium.").font = Font(
        name=FONT, italic=True, color="808080")


def build_policies(wb):
    ws = wb.create_sheet("Policies")
    cols = POLICY_COLS
    write_headers(ws, cols)
    style_body(ws, cols)
    C = {name: col(cols, name) for name, *_ in cols}
    freq_ref = list_ref("Frequency")
    ppy_ref = list_ref("Payments / Year")
    for r in range(2, ROWS + 2):
        ws[f"{C['Annual Premium']}{r}"] = (
            f'=IF(OR({C["Premium"]}{r}="",{C["Frequency"]}{r}=""),"",'
            f'IF({C["Frequency"]}{r}="Single Pay",0,'
            f'{C["Premium"]}{r}*IFERROR(INDEX({ppy_ref},MATCH({C["Frequency"]}{r},{freq_ref},0)),0)))'
        )
        ws[f"{C['Term Ends']}{r}"] = (
            f'=IF(OR({C["Issue Date"]}{r}="",{C["Term Length (Yrs)"]}{r}=""),"",'
            f'EDATE({C["Issue Date"]}{r},12*{C["Term Length (Yrs)"]}{r}))'
        )
        ws[f"{C['Days Until Due']}{r}"] = (
            f'=IF({C["Next Premium Due"]}{r}="","",{C["Next Premium Due"]}{r}-TODAY())'
        )
        ws[f"{C['Total Paid']}{r}"] = (
            f'=IF({C["Policy Number"]}{r}="","",'
            f"SUMIFS(Payments!$C$2:$C${ROWS + 1},Payments!$B$2:$B${ROWS + 1},{C['Policy Number']}{r}))"
        )
        ws[f"{C['Primary Benef. %']}{r}"] = (
            f'=IF({C["Policy Number"]}{r}="","",'
            f'SUMIFS(Beneficiaries!$E$2:$E${ROWS + 1},Beneficiaries!$A$2:$A${ROWS + 1},{C["Policy Number"]}{r},'
            f'Beneficiaries!$D$2:$D${ROWS + 1},"Primary"))'
        )

    add_dropdown(ws, cols, "Policy Type", "Policy Type")
    add_dropdown(ws, cols, "Status", "Status")
    add_dropdown(ws, cols, "Frequency", "Frequency")
    add_dropdown(ws, cols, "Payment Method", "Payment Method")
    day = DataValidation(type="whole", operator="between", formula1="1", formula2="31", allow_blank=True)
    day.error = "Enter a day of the month from 1 to 31."
    ws.add_data_validation(day)
    d = C["Draft Day of Month"]
    day.add(f"{d}2:{d}{ROWS + 1}")

    # Highlights: due within 14 days (amber), overdue (red), beneficiaries not 100% (red text)
    due = C["Days Until Due"]
    rng = f"{due}2:{due}{ROWS + 1}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ISNUMBER({due}2),{due}2<0)'],
        fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=FONT, bold=True, color="9C0006")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ISNUMBER({due}2),{due}2>=0,{due}2<=14)'],
        fill=PatternFill("solid", fgColor="FFE699")))
    pb = C["Primary Benef. %"]
    ws.conditional_formatting.add(f"{pb}2:{pb}{ROWS + 1}", FormulaRule(
        formula=[f'AND(ISNUMBER({pb}2),ROUND({pb}2,4)<>1)'],
        font=Font(name=FONT, bold=True, color="C00000")))
    st = C["Status"]
    ws.conditional_formatting.add(f"{st}2:{st}{ROWS + 1}", FormulaRule(
        formula=[f'OR({st}2="Lapsed",{st}2="Declined",{st}2="Cancelled")'],
        font=Font(name=FONT, bold=True, color="C00000")))
    ws.conditional_formatting.add(f"{st}2:{st}{ROWS + 1}", FormulaRule(
        formula=[f'{st}2="In Force"'], font=Font(name=FONT, bold=True, color="006100")))

    # Example row (row 2) so the expected format is clear; delete or overwrite it.
    today = dt.date.today()
    next_draft = today.replace(day=15)
    if next_draft <= today:
        next_draft = (today.replace(day=28) + dt.timedelta(days=4)).replace(day=15)
    example = {
        "Policy Number": "EX-0001",
        "Insured": "Jane Example",
        "Owner": "Jane Example",
        "Carrier": "Example Life Co.",
        "Policy Type": "Term",
        "Status": "In Force",
        "Face Amount": 500000,
        "Premium": 42.50,
        "Frequency": "Monthly",
        "Application Date": dt.date(today.year - 1, 3, 1),
        "Issue Date": dt.date(today.year - 1, 4, 15),
        "Term Length (Yrs)": 20,
        "Payment Method": "Bank Draft",
        "Draft Day of Month": 15,
        "Next Premium Due": next_draft,
        "Agent / Contact": "Sam Agent (555) 010-0000",
        "Notes": "EXAMPLE ROW - overwrite or delete",
    }
    for name, value in example.items():
        cell = ws[f"{C[name]}2"]
        cell.value = value
        cell.font = EXAMPLE_FONT
    return C


def build_beneficiaries(wb, pol):
    ws = wb.create_sheet("Beneficiaries")
    cols = BENEF_COLS
    write_headers(ws, cols)
    style_body(ws, cols)
    for r in range(2, ROWS + 2):
        ws[f"F{r}"] = (
            f'=IF(A{r}="","",IFERROR(INDEX(Policies!${pol["Insured"]}$2:${pol["Insured"]}${ROWS + 1},'
            f'MATCH(A{r},Policies!${pol["Policy Number"]}$2:${pol["Policy Number"]}${ROWS + 1},0)),"Policy # not found"))'
        )
    add_dropdown(ws, cols, "Relationship", "Relationship")
    add_dropdown(ws, cols, "Type", "Beneficiary Type")
    pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    pct.error = "Enter a percentage from 0% to 100%."
    ws.add_data_validation(pct)
    pct.add(f"E2:E{ROWS + 1}")
    for c, v in zip("ABCDEG", ["EX-0001", "John Example", "Spouse", "Primary", 1, "EXAMPLE ROW"]):
        ws[f"{c}2"] = v
        ws[f"{c}2"].font = EXAMPLE_FONT
    for c, v in zip("ABCDEG", ["EX-0001", "Kid Example", "Child", "Contingent", 1, "EXAMPLE ROW"]):
        ws[f"{c}3"] = v
        ws[f"{c}3"].font = EXAMPLE_FONT


def build_payments(wb, pol):
    ws = wb.create_sheet("Payments")
    cols = PAYMENT_COLS
    write_headers(ws, cols)
    style_body(ws, cols)
    for r in range(2, ROWS + 2):
        ws[f"E{r}"] = (
            f'=IF(B{r}="","",IFERROR(INDEX(Policies!${pol["Insured"]}$2:${pol["Insured"]}${ROWS + 1},'
            f'MATCH(B{r},Policies!${pol["Policy Number"]}$2:${pol["Policy Number"]}${ROWS + 1},0)),"Policy # not found"))'
        )
    add_dropdown(ws, cols, "Method", "Payment Method")
    today = dt.date.today()
    first = today.replace(day=1)
    prev = (first - dt.timedelta(days=1)).replace(day=15)
    for row, d in ((2, prev), (3, first.replace(day=15) if today.day >= 15 else None)):
        if d is None:
            continue
        for c, v in zip("ABCDF", [d, "EX-0001", 42.50, "Bank Draft", "EXAMPLE ROW"]):
            ws[f"{c}{row}"] = v
            ws[f"{c}{row}"].font = EXAMPLE_FONT


def build_summary(wb, pol):
    ws = wb.create_sheet("Summary", 1)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    title = ws["A1"]
    title.value = "Life Insurance Tracker - Summary"
    title.font = Font(name=FONT, bold=True, size=14, color="1F3864")
    ws["A2"] = '="As of "&TEXT(TODAY(),"mmm d, yyyy")'
    ws["A2"].font = Font(name=FONT, italic=True, color="808080")

    P = lambda name: f"Policies!${pol[name]}$2:${pol[name]}${ROWS + 1}"
    status, face, annual = P("Status"), P("Face Amount"), P("Annual Premium")
    due, polno = P("Days Until Due"), P("Policy Number")
    pb = P("Primary Benef. %")

    kpis = [
        ("Policies tracked", f'=COUNTA({polno})', "0"),
        ("Policies in force", f'=COUNTIFS({status},"In Force")', "0"),
        ("Total coverage in force", f'=SUMIFS({face},{status},"In Force")', CURRENCY0),
        ("Annual premium (in force)", f'=SUMIFS({annual},{status},"In Force")', CURRENCY),
        ("Monthly equivalent", "=B8/12", CURRENCY),
        ("Premiums overdue", f'=COUNTIFS({due},"<0",{status},"In Force")', "0"),
        ("Premiums due in next 14 days", f'=COUNTIFS({due},">=0",{due},"<=14",{status},"In Force")', "0"),
        ("Policies w/ primary beneficiaries <> 100%",
         f'=SUMPRODUCT(({polno}<>"")*ISNUMBER({pb})*(ROUND(N(+{pb}),4)<>1))', "0"),
        ("Total premiums paid (all time)", f"=SUM(Payments!$C$2:$C${ROWS + 1})", CURRENCY),
    ]
    ws["A4"] = "Key figures"
    ws["A4"].font = Font(name=FONT, bold=True, color="FFFFFF")
    ws["A4"].fill = HEADER_FILL
    ws["B4"].fill = HEADER_FILL
    for i, (label, formula, fmt) in enumerate(kpis, start=5):
        ws[f"A{i}"] = label
        ws[f"A{i}"].font = Font(name=FONT)
        ws[f"B{i}"] = formula
        ws[f"B{i}"].number_format = fmt
        ws[f"B{i}"].font = Font(name=FONT, bold=True)
        for c in "AB":
            ws[f"{c}{i}"].border = BORDER

    start = 5 + len(kpis) + 2
    for c, h in zip("ABCD", ["By status", "Policies", "Face amount", "Annual premium"]):
        cell = ws[f"{c}{start}"]
        cell.value = h
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for i, s in enumerate(LISTS["Status"], start=start + 1):
        ws[f"A{i}"] = s
        ws[f"B{i}"] = f'=COUNTIFS({status},A{i})'
        ws[f"C{i}"] = f'=SUMIFS({face},{status},A{i})'
        ws[f"D{i}"] = f'=SUMIFS({annual},{status},A{i})'
        ws[f"B{i}"].number_format = "0"
        ws[f"C{i}"].number_format = CURRENCY0
        ws[f"D{i}"].number_format = CURRENCY
        for c in "ABCD":
            ws[f"{c}{i}"].border = BORDER
            ws[f"{c}{i}"].font = Font(name=FONT)

    start = start + len(LISTS["Status"]) + 3
    for c, h in zip("ABCD", ["By policy type", "Policies", "Face amount", "Annual premium"]):
        cell = ws[f"{c}{start}"]
        cell.value = h
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    ptype = P("Policy Type")
    for i, t in enumerate(LISTS["Policy Type"], start=start + 1):
        ws[f"A{i}"] = t
        ws[f"B{i}"] = f'=COUNTIFS({ptype},A{i})'
        ws[f"C{i}"] = f'=SUMIFS({face},{ptype},A{i})'
        ws[f"D{i}"] = f'=SUMIFS({annual},{ptype},A{i})'
        ws[f"B{i}"].number_format = "0"
        ws[f"C{i}"].number_format = CURRENCY0
        ws[f"D{i}"].number_format = CURRENCY
        for c in "ABCD":
            ws[f"{c}{i}"].border = BORDER
            ws[f"{c}{i}"].font = Font(name=FONT)


def build_howto(wb):
    ws = wb.active
    ws.title = "How to Use"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Life Insurance Tracker", Font(name=FONT, bold=True, size=16, color="1F3864")),
        ("", None),
        ("What goes where", Font(name=FONT, bold=True, size=12)),
        ("Policies - one row per policy. Policy Number is the key that links everything together.", None),
        ("Beneficiaries - one row per beneficiary. Use the same Policy Number as on the Policies sheet.", None),
        ("Payments - log each premium payment with its Policy Number; totals roll up to Policies automatically.", None),
        ("Summary - read-only dashboard: coverage in force, premiums, what's due or overdue.", None),
        ("Lists - the dropdown options. Add or rename items here (e.g. new carriers' product types or statuses).", None),
        ("", None),
        ("Color legend", Font(name=FONT, bold=True, size=12)),
        ("Cream cells with blue text = you type here.", None),
        ("Grey cells = calculated automatically. Don't type over them.", None),
        ("Blue italic rows marked EXAMPLE ROW show the expected format - overwrite or delete them.", None),
        ("", None),
        ("Highlights on the Policies sheet", Font(name=FONT, bold=True, size=12)),
        ("Days Until Due: red = overdue, amber = due within 14 days. Update Next Premium Due after each payment.", None),
        ("Primary Benef. %: red text when primary beneficiary shares for that policy don't add up to 100%.", None),
        ("Status: green = In Force; red = Lapsed / Cancelled / Declined.", None),
        ("", None),
        ("Tips", Font(name=FONT, bold=True, size=12)),
        ("Use the filter arrows in the header row to sort or filter (e.g. show only In Force).", None),
        ("Term Ends is calculated from Issue Date + Term Length (Yrs); leave Term Length blank for permanent policies.", None),
        ("Don't store full bank account or Social Security numbers in this file.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=2, value=text)
        cell.font = font or Font(name=FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=11, column=2).fill = INPUT_FILL
    ws.cell(row=11, column=2).font = INPUT_FONT
    ws.cell(row=12, column=2).fill = CALC_FILL
    ws.cell(row=13, column=2).font = EXAMPLE_FONT


def build(output):
    wb = Workbook()
    build_howto(wb)
    pol = build_policies(wb)
    build_beneficiaries(wb, pol)
    build_payments(wb, pol)
    build_lists(wb)
    build_summary(wb, pol)
    wb.calculation.fullCalcOnLoad = True  # no cached values from openpyxl; recalc on open
    wb.save(output)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--output", default="life_insurance_tracker.xlsx")
    args = ap.parse_args()
    build(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
